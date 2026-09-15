"""Geometry and text-fit audit for a generated deck.

Nothing here renders PowerPoint, so it cannot prove typography. It does prove
the things the generator controls: every shape sits inside the slide, no text
box collides with another, and no run of text is obviously too long for the
box it was given.
"""
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_IN = 914400.0
CHAR_W = {  # rough advance width as a fraction of point size, Calibri-ish
    True: 0.52,   # bold
    False: 0.49,
}


def boxes(slide):
    out = []
    for sh in slide.shapes:
        if sh.left is None or sh.top is None:
            continue
        out.append(sh)
    return out


def text_of(sh):
    if not sh.has_text_frame:
        return "", 0, False
    tf = sh.text_frame
    longest = 0.0
    bold = False
    txt = []
    for p in tf.paragraphs:
        for r in p.runs:
            txt.append(r.text)
            sz = r.font.size.pt if r.font.size else 18
            longest = max(longest, sz)
            bold = bold or bool(r.font.bold)
    return "".join(txt), longest, bold


def audit(path, slide_w=13.333, slide_h=7.5):
    prs = Presentation(path)
    problems = []
    for i, slide in enumerate(prs.slides, 1):
        shs = boxes(slide)
        for sh in shs:
            L, T = sh.left / EMU_IN, sh.top / EMU_IN
            W = (sh.width or 0) / EMU_IN
            H = (sh.height or 0) / EMU_IN
            if L < -0.01 or T < -0.01 or L + W > slide_w + 0.01 or T + H > slide_h + 0.01:
                problems.append(
                    f"slide {i}: shape off-slide  ({L:.2f},{T:.2f}) {W:.2f}x{H:.2f}"
                    f"  right={L+W:.2f} bottom={T+H:.2f}  text={text_of(sh)[0][:44]!r}")
            if T + H > 6.88 and not (T >= 6.90 and H < 0.40):
                problems.append(
                    f"slide {i}: intrudes on the footer band  bottom={T+H:.2f}"
                    f"  text={text_of(sh)[0][:40]!r}")
            body, sz, bold = text_of(sh)
            if body and W > 0.3 and H > 0.1 and sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                per_line = max(1, int(W * 72 / (sz * CHAR_W[bold])))
                lines = 0
                for para in body.split("\n"):
                    lines += max(1, -(-len(para) // per_line))
                needed = lines * sz * 1.30 / 72.0
                if needed > H + 0.30:
                    problems.append(
                        f"slide {i}: text may overflow  box {W:.2f}x{H:.2f}in @{sz}pt needs "
                        f"~{needed:.2f}in ({lines} lines)  {body[:52]!r}")
        # shape-vs-shape: text that collides with other text reads as a bug
        tb = []
        for sh in shs:
            if sh.has_text_frame and sh.text_frame.text.strip() and sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                tb.append((sh.left / EMU_IN, sh.top / EMU_IN,
                           (sh.width or 0) / EMU_IN, (sh.height or 0) / EMU_IN,
                           text_of(sh)[0]))
        for a_i in range(len(tb)):
            for b_i in range(a_i + 1, len(tb)):
                ax, ay, aw, ah, at = tb[a_i]
                bx, by, bw, bh, bt = tb[b_i]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.06 and oy > 0.06:
                    problems.append(
                        f"slide {i}: text boxes overlap by {ox:.2f}x{oy:.2f}in  "
                        f"{at[:30]!r} / {bt[:30]!r}")
        # a text box sitting on top of a card (but not inside it) reads as broken
        cards = []
        for sh in shs:
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                W_ = (sh.width or 0) / EMU_IN
                H_ = (sh.height or 0) / EMU_IN
                if W_ >= 1.5 and H_ >= 1.0:
                    cards.append((sh.left / EMU_IN, sh.top / EMU_IN, W_, H_))
        for ax, ay, aw, ah, at in tb:
            for cx, cy, cw_, ch in cards:
                inside = (ax >= cx - 0.05 and ay >= cy - 0.05
                          and ax + aw <= cx + cw_ + 0.05 and ay + ah <= cy + ch + 0.05)
                if inside:
                    continue
                ox = min(ax + aw, cx + cw_) - max(ax, cx)
                oy = min(ay + ah, cy + ch) - max(ay, cy)
                if ox > 0.10 and oy > 0.10:
                    problems.append(
                        f"slide {i}: text sits on a card by {ox:.2f}x{oy:.2f}in  {at[:44]!r}")
                    break
    return prs, problems


if __name__ == "__main__":
    for path in sys.argv[1:]:
        prs, probs = audit(path)
        n = len(prs.slides._sldIdLst)
        print(f"\n=== {path}  ({n} slides) ===")
        if not probs:
            print("  clean — every shape inside the slide, no obvious text overflow")
        for p in probs:
            print("  " + p)
