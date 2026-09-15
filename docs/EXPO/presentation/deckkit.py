"""Shared slide primitives for the expo decks.

Authoring tooling, not demo runtime. Renders real .pptx so the deck opens on
any expo machine (PowerPoint, Keynote, Google Slides) without this repo.
Palette and voice mirror the Control Centre itself.
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- palette (the Control Centre's own) ---------------------------------
RED = RGBColor(0xD8, 0x23, 0x2A)
RED_DEEP = RGBColor(0xA8, 0x16, 0x1C)
INK = RGBColor(0x2B, 0x2B, 0x2B)
INK_SOFT = RGBColor(0x4A, 0x45, 0x40)
MUTED = RGBColor(0x7A, 0x74, 0x69)
LINE = RGBColor(0xD9, 0xD2, 0xC7)
PAPER = RGBColor(0xFA, 0xF7, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OK = RGBColor(0x1F, 0x7A, 0x4D)
OK_BG = RGBColor(0xE6, 0xF4, 0xEC)
OK_LN = RGBColor(0xB8, 0xDD, 0xC8)
WARN = RGBColor(0x8A, 0x64, 0x10)
WARN_BG = RGBColor(0xFD, 0xF3, 0xE0)
WARN_LN = RGBColor(0xE6, 0xCF, 0x9A)
INFO = RGBColor(0x1D, 0x5C, 0x7A)
INFO_BG = RGBColor(0xE8, 0xF1, 0xF5)
INFO_LN = RGBColor(0xB6, 0xD4, 0xE0)
SLATE = RGBColor(0x4C, 0x55, 0x60)
SLATE_BG = RGBColor(0xEE, 0xF0, 0xF2)
SLATE_LN = RGBColor(0xCC, 0xD3, 0xDA)

FONT = "Calibri"          # present wherever PowerPoint is; survives any machine
MONO = "Consolas"

W, H = 13.333, 7.5
ML, MR = 0.78, 0.78
CW = W - ML - MR          # content width


def deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def _noshadow(shape):
    shape.shadow.inherit = False


# --- v2: accent family, icon tiles and glass-style surfaces ---------------
ICON_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "icons")

ACCENTS = {
    "red":    (RED,                        RGBColor(0xFD, 0xEC, 0xEC), RGBColor(0xF4, 0xC3, 0xC4)),
    "teal":   (RGBColor(0x0F, 0x7A, 0x6C), RGBColor(0xE4, 0xF3, 0xF0), RGBColor(0xB3, 0xDD, 0xD5)),
    "amber":  (RGBColor(0xB0, 0x77, 0x0C), RGBColor(0xFD, 0xF1, 0xDD), RGBColor(0xEA, 0xCF, 0x9B)),
    "violet": (RGBColor(0x6B, 0x4F, 0xA8), RGBColor(0xF0, 0xEC, 0xF9), RGBColor(0xCF, 0xC2, 0xE8)),
    "blue":   (RGBColor(0x1D, 0x5C, 0x7A), RGBColor(0xE6, 0xF1, 0xF6), RGBColor(0xB4, 0xD3, 0xE0)),
    "green":  (RGBColor(0x1F, 0x7A, 0x4D), RGBColor(0xE6, 0xF4, 0xEC), RGBColor(0xB6, 0xDC, 0xC7)),
    "slate":  (RGBColor(0x4C, 0x55, 0x60), RGBColor(0xEE, 0xF0, 0xF2), RGBColor(0xCC, 0xD3, 0xDA)),
}


def glassy(shape, accent=None):
    """A light top-down gradient and a bright edge — the closest a static
    slide gets to the films' Liquid Glass surfaces."""
    fill = shape.fill
    fill.gradient()
    fill.gradient_angle = 90.0
    stops = fill.gradient_stops
    if accent:
        _, bg, _ = ACCENTS[accent]
        stops[0].color.rgb = WHITE
        stops[1].color.rgb = bg
    else:
        stops[0].color.rgb = WHITE
        stops[1].color.rgb = RGBColor(0xF6, 0xF2, 0xEC)
    stops[0].position = 0.0
    stops[1].position = 1.0
    shape.line.color.rgb = ACCENTS[accent][2] if accent else LINE
    shape.line.width = Pt(1.0)
    _noshadow(shape)
    return shape


def icon_tile(s, x, y, size, name, accent="slate"):
    """Rounded tinted square with the film's own icon centred in it."""
    fg, bg, ln = ACCENTS[accent]
    t = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(size), Inches(size))
    t.adjustments[0] = 0.30
    t.fill.solid(); t.fill.fore_color.rgb = bg
    t.line.color.rgb = ln; t.line.width = Pt(1.0)
    _noshadow(t)
    png = os.path.join(ICON_DIR, "%s-%s.png" % (name, accent))
    if not os.path.exists(png):
        raise ValueError("no icon %r in accent %r — run render_icons.js" % (name, accent))
    ins = size * 0.52
    s.shapes.add_picture(png, Inches(x + (size - ins) / 2), Inches(y + (size - ins) / 2),
                         width=Inches(ins), height=Inches(ins))
    return t


def slide(prs, chip=None, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Pt(5))
    bar.fill.solid(); bar.fill.fore_color.rgb = RED; bar.line.fill.background()
    _noshadow(bar)
    if chip is not None:
        footer(s, chip)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def footer(s, chip):
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ML), Inches(6.95), Inches(0.26), Inches(0.26))
    dot.fill.solid(); dot.fill.fore_color.rgb = RED; dot.line.fill.background(); _noshadow(dot)
    t = dot.text_frame; t.margin_left = t.margin_right = t.margin_top = t.margin_bottom = 0
    p = t.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "MS"
    r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT

    text(s, "MAPLESURE INSURANCE  ·  CONTROL CENTER", ML + 0.38, 6.96, 5.4, 0.26,
         size=9.5, color=MUTED, bold=True, space=1.6)
    text(s, chip, W - MR - 5.4, 6.96, 5.4, 0.26,
         size=9.5, color=MUTED, bold=True, space=1.6, align=PP_ALIGN.RIGHT)


def text(s, body, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         space=0, line=1.18, font=FONT, anchor=MSO_ANCHOR.TOP, italic=False):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = body.split("\n") if isinstance(body, str) else list(body)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        _rich(p, ln, size, color, bold, space, font, italic)
    return box


def _rich(p, ln, size, color, bold, space, font, italic):
    """**bold** and {accent} inline markers keep the call sites readable."""
    import re
    for part in re.split(r"(\*\*.+?\*\*|\{.+?\})", ln):
        if not part:
            continue
        r = p.add_run()
        if part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]; r.font.bold = True; r.font.color.rgb = color
        elif part.startswith("{") and part.endswith("}"):
            r.text = part[1:-1]; r.font.bold = True; r.font.color.rgb = RED
        else:
            r.text = part; r.font.bold = bold; r.font.color.rgb = color
        r.font.size = Pt(size); r.font.name = font; r.font.italic = italic
        if space:
            r.font._rPr.set("spc", str(int(space * 100)))


def kicker(s, body, y=0.60):
    text(s, body.upper(), ML, y, CW, 0.3, size=11.5, color=RED, bold=True, space=2.2)


def title(s, body, y=0.92, size=33):
    text(s, body, ML, y, CW, 1.25, size=size, color=INK, bold=True, line=1.06)


def lede(s, body, y, size=15.5, w=None, color=MUTED):
    text(s, body, ML, y, w or CW, 1.0, size=size, color=color, line=1.32)



def _head_h(body, w_in, size, bold=True, line=1.14):
    """Estimated rendered height of a heading, so the body beneath it can never
    be overlapped by a heading that wrapped."""
    import math
    per_line = max(1, int(w_in * 72 / (size * (0.545 if bold else 0.515))))
    lines = 0
    for para in str(body).split("\n"):
        lines += max(1, math.ceil(len(para) / per_line))
    return lines * size * line / 72.0


def card(s, x, y, w, h, head, body, num=None, accent=RED, fill=WHITE, icon=None, tone=None):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.070
    if tone:
        glassy(box, tone)
    else:
        box.fill.solid(); box.fill.fore_color.rgb = fill
        box.line.color.rgb = LINE; box.line.width = Pt(1.1)
    _noshadow(box)
    box.text_frame.text = ""
    cy = y + 0.26
    head_w = w - 0.6
    if icon:
        # top-right, so an icon never costs the card any height
        icon_tile(s, x + w - 0.30 - 0.46, y + 0.26, 0.46, icon, tone or "red")
        head_w = w - 1.16
    if num:
        text(s, num, x + 0.3, cy, w - 0.6, 0.24, size=10.5,
             color=(ACCENTS[tone][0] if tone else accent), bold=True, space=2.0)
        cy += 0.34
    hh = _head_h(head, head_w, 16.5, True, 1.1)
    text(s, head, x + 0.3, cy, head_w, hh + 0.05, size=16.5, color=INK, bold=True, line=1.1)
    by = cy + hh + 0.19
    avail = h - (by - y) - 0.24
    # Fit the body to the card by stepping the size down a little rather than
    # letting it run past the border. Raise only if even the floor overflows,
    # so a card that is genuinely too small fails the build instead of shipping.
    for bsize in (12.5, 12.0, 11.5, 11.0, 10.5):
        need = _head_h(body, w - 0.6, bsize, False, 1.30)
        if need <= avail + 0.02:
            break
    else:
        raise ValueError(
            "card() body overflows even at 10.5pt: %r needs %.2fin but the card "
            "leaves %.2fin. Raise h to about %.2f or shorten the body."
            % (head[:36], need, avail, h + (need - avail)))
    text(s, body, x + 0.3, by, w - 0.6, avail, size=bsize, color=MUTED, line=1.30)
    return box


def tile(s, x, y, w, h, value, label, value_color=INK, icon=None, tone=None):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.10
    if tone:
        glassy(box, tone)
    else:
        box.fill.solid(); box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = LINE; box.line.width = Pt(1.1)
    _noshadow(box)
    vw = w - 0.5
    if icon:
        icon_tile(s, x + w - 0.26 - 0.40, y + 0.20, 0.40, icon, tone or "slate")
        vw = w - 0.92
    text(s, value, x + 0.26, y + 0.20, vw, 0.6, size=31, color=value_color, bold=True, line=1.0)
    text(s, label, x + 0.26, y + 0.82, w - 0.5, h - 1.0, size=11.5, color=MUTED, line=1.24)
    return box


def badge(s, x, y, label, fg, bg, ln, w=1.32, h=0.33, size=10.5):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.adjustments[0] = 0.5
    b.fill.solid(); b.fill.fore_color.rgb = bg
    b.line.color.rgb = ln; b.line.width = Pt(1.1)
    _noshadow(b)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label.upper()
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = fg; r.font.name = FONT
    r.font._rPr.set("spc", "120")
    return b


def pill(s, x, y, label, w=None, size=11.5):
    w = w or (0.20 + 0.088 * len(label))
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.36))
    b.adjustments[0] = 0.5
    b.fill.solid(); b.fill.fore_color.rgb = WHITE
    b.line.color.rgb = LINE; b.line.width = Pt(1.1)
    _noshadow(b)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.color.rgb = INK_SOFT; r.font.name = FONT
    return x + w


def quote(s, x, y, w, h, head, body):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.045
    box.fill.solid(); box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = LINE; box.line.width = Pt(1.1)
    _noshadow(box)
    edge = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.06), Pt(4.5), Inches(h - 0.12))
    edge.fill.solid(); edge.fill.fore_color.rgb = RED; edge.line.fill.background(); _noshadow(edge)
    hh = _head_h(head, w - 0.66, 16, True, 1.14)
    text(s, head, x + 0.34, y + 0.24, w - 0.66, hh + 0.05, size=16, color=INK, bold=True, line=1.14)
    if body:
        by = y + 0.24 + hh + 0.17
        bh = h - (by - y) - 0.22
        if bh < 0.22:
            raise ValueError(
                "quote() body has no room: head %r wraps to %.2fin inside a %.2fin box. "
                "Shorten the head or raise h." % (head[:40], hh, h))
        text(s, body, x + 0.34, by, w - 0.66, bh, size=12.5, color=MUTED, line=1.30)
    return box


def rows(s, x, y, w, items, label_w=2.3, size=12.5, rh=0.46, label_color=INK):
    """Definition rows with a hairline between them.

    Each row is as tall as its own content, so a long value can never run into
    the row beneath it."""
    cy = y
    vw = w - label_w - 0.2
    for i, (k, v) in enumerate(items):
        hk = _head_h(k, label_w, size, True, 1.2)
        hv = _head_h(v, vw, size, False, 1.24)
        rowh = max(rh, hk, hv)
        text(s, k, x, cy + 0.04, label_w, hk + 0.06, size=size, color=label_color, bold=True, line=1.2)
        text(s, v, x + label_w + 0.2, cy + 0.04, vw, hv + 0.06, size=size, color=MUTED, line=1.24)
        if i < len(items) - 1:
            ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(cy + rowh + 0.06),
                                    Inches(w), Pt(0.85))
            ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); _noshadow(ln)
        cy += rowh + 0.20
    return cy


def table(s, x, y, w, headers, data, col_w, size=11.5, rh=0.4):
    """Plain-drawn table — python-pptx tables fight the theme, shapes do not."""
    cx = x
    for i, hcell in enumerate(headers):
        text(s, hcell.upper(), cx, y, col_w[i], 0.3, size=9.5, color=MUTED, bold=True, space=1.4)
        cx += col_w[i]
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.32), Inches(w), Pt(1.2))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); _noshadow(ln)
    cy = y + 0.46
    for r_i, row in enumerate(data):
        cx = x
        rowh = rh
        for i, cell in enumerate(row):
            rowh = max(rowh, _head_h(cell, col_w[i] - 0.14, size, i == 0, 1.2) + 0.10)
        for i, cell in enumerate(row):
            col = INK if i == 0 else MUTED
            text(s, cell, cx, cy, col_w[i] - 0.14, rowh, size=size, color=col, line=1.2)
            cx += col_w[i]
        cy += rowh
        if r_i < len(data) - 1:
            ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(cy - 0.06), Inches(w), Pt(0.85))
            ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); _noshadow(ln)
    return cy


def flow(s, y, nodes, d=0.80):
    """nodes = [(label, sub, kind)] with kind in {'step','gate'}.

    Laid out on an even pitch across the content width so the labels get the
    whole cell and cannot collide with their neighbours.
    """
    n = len(nodes)
    pitch = CW / n
    for i, (label, sub, kind) in enumerate(nodes):
        is_gate = kind == "gate"
        fg, bg, ln = (RED, RGBColor(0xFD, 0xEC, 0xED), RED) if is_gate else (OK, OK_BG, OK_LN)
        cx = ML + pitch * i + pitch / 2
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(y), Inches(d), Inches(d))
        o.fill.solid(); o.fill.fore_color.rgb = bg
        o.line.color.rgb = ln; o.line.width = Pt(2.25); _noshadow(o)
        tf = o.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pr = tf.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER
        r = pr.add_run(); r.text = ("G%d" % (sum(1 for k in nodes[:i + 1] if k[2] == "gate"))) if is_gate else str(i + 1)
        r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = fg; r.font.name = FONT
        text(s, label, ML + pitch * i + 0.04, y + d + 0.15, pitch - 0.08, 0.28,
             size=12, color=INK_SOFT, bold=True, align=PP_ALIGN.CENTER, line=1.1)
        text(s, sub, ML + pitch * i + 0.04, y + d + 0.43, pitch - 0.08, 0.46,
             size=9.5, color=MUTED, align=PP_ALIGN.CENTER, line=1.14)
        if i < n - 1:
            c = RED if (is_gate or nodes[i + 1][2] == "gate") else OK
            x0 = cx + d / 2 + 0.10
            x1 = cx + pitch - d / 2 - 0.10
            bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(y + d / 2 - 0.012),
                                     Inches(x1 - x0), Pt(2.0))
            bar.fill.solid(); bar.fill.fore_color.rgb = c; bar.line.fill.background(); _noshadow(bar)


def shot(s, path, x, y, w, caption=None):
    """Screenshot in a light frame, optional caption beneath."""
    frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(x - 0.045), Inches(y - 0.045),
                               Inches(w + 0.09), Inches(w * 784 / 1547 + 0.09))
    frame.adjustments[0] = 0.02
    frame.fill.solid(); frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = LINE; frame.line.width = Pt(1.1); _noshadow(frame)
    s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    h = w * 784 / 1547
    if caption:
        text(s, caption, x, y + h + 0.16, w, 0.4, size=10.5, color=MUTED, line=1.24)
    return y + h


def movie(s, mp4, poster, x, y, w, h):
    return s.shapes.add_movie(mp4, Inches(x), Inches(y), Inches(w), Inches(h),
                              poster_frame_image=poster, mime_type="video/mp4")


# --- v2 components mirroring the films -----------------------------------

def persona(s, x, y, w, h, initials, name, role, tone, ctx, out_html, out_bad=None):
    """A developer and whatever context they happened to have."""
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.065
    glassy(box, tone)
    fg, bg, ln = ACCENTS[tone]

    av = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.28), Inches(y + 0.26), Inches(0.52), Inches(0.52))
    av.fill.solid(); av.fill.fore_color.rgb = bg
    av.line.color.rgb = ln; av.line.width = Pt(1.6)
    _noshadow(av)
    tf = av.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pr = tf.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER
    r = pr.add_run(); r.text = initials
    r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = fg; r.font.name = FONT

    text(s, name, x + 0.92, y + 0.30, w - 1.2, 0.26, size=16, color=INK, bold=True)
    text(s, role, x + 0.92, y + 0.56, w - 1.2, 0.24, size=11.5, color=MUTED)

    cy = y + 0.92
    limit = y + h - 0.86          # never let a context line reach the output box
    for line in ctx:
        hh = _head_h(line, w - 0.86, 11.5, False, 1.24)
        if cy + hh > limit:
            raise ValueError(
                "persona() context overflows for %r — raise h to about %.2f or shorten a line"
                % (name, h + (cy + hh - limit)))
        text(s, "·", x + 0.30, cy, 0.18, 0.26, size=12.5, color=MUTED, bold=True)
        text(s, line, x + 0.52, cy, w - 0.86, hh + 0.06, size=11.5, color=INK_SOFT, line=1.24)
        cy += max(0.28, hh + 0.08)

    ob = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.28), Inches(y + h - 0.70),
                            Inches(w - 0.56), Inches(0.48))
    ob.adjustments[0] = 0.22
    ob.fill.solid(); ob.fill.fore_color.rgb = RGBColor(0x1B, 0x1F, 0x24)
    ob.line.fill.background(); _noshadow(ob)
    tf = ob.text_frame
    tf.margin_left = Inches(0.16); tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = out_html
    r.font.size = Pt(11.5); r.font.name = MONO; r.font.color.rgb = RGBColor(0xC9, 0xD3, 0xDE)
    if out_bad:
        r2 = p.add_run(); r2.text = out_bad
        r2.font.size = Pt(11.5); r2.font.name = MONO; r2.font.color.rgb = RGBColor(0xE0, 0x60, 0x5E)
    return box


def artifact_chip(s, x, y, w, icon, tone, name, meta):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.66))
    box.adjustments[0] = 0.20
    glassy(box, tone)
    icon_tile(s, x + 0.13, y + 0.13, 0.40, icon, tone)
    text(s, name, x + 0.66, y + 0.12, w - 0.78, 0.24, size=12.5, color=INK, bold=True)
    text(s, meta, x + 0.66, y + 0.36, w - 0.78, 0.22, size=10.5, color=MUTED)
    return box


def lane(s, y, steps, tone=None, step_w=2.18, gap=0.28, h=2.18, compact=False):
    """A row of governed steps, each naming the artifact it leaves behind.
    steps = [(icon, accent, title, sub, out)]"""
    n = len(steps)
    total = n * step_w + (n - 1) * gap
    x = (W - total) / 2
    for i, (icon, accent, ttl, sub, out) in enumerate(steps):
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(step_w), Inches(h))
        box.adjustments[0] = 0.085
        glassy(box, accent)
        icon_tile(s, x + (step_w - 0.52) / 2, y + 0.18, 0.52, icon, accent)
        text(s, ttl, x + 0.12, y + 0.80, step_w - 0.24, 0.28, size=14.5, color=INK, bold=True,
             align=PP_ALIGN.CENTER)
        if not compact:
            text(s, sub, x + 0.14, y + 1.10, step_w - 0.28, 0.52, size=10.5, color=MUTED,
                 align=PP_ALIGN.CENTER, line=1.24)
        ow = 0.24 + 0.075 * len(out)
        ob = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + (step_w - ow) / 2),
                                Inches(y + h - 0.44), Inches(ow), Inches(0.30))
        ob.adjustments[0] = 0.5
        fg, bg, ln = ACCENTS[accent]
        ob.fill.solid(); ob.fill.fore_color.rgb = bg
        ob.line.color.rgb = ln; ob.line.width = Pt(1.0)
        _noshadow(ob)
        tf = ob.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = out
        r.font.size = Pt(9.5); r.font.name = MONO; r.font.color.rgb = fg; r.font.bold = True

        if i < n - 1:
            nxt = steps[i + 1][1]
            c = ACCENTS["red"][0] if (accent == "red" or nxt == "red") else ACCENTS["green"][0]
            aw = max(0.26, gap - 0.08)
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + step_w + (gap - aw) / 2),
                                    Inches(y + h / 2 - 0.085), Inches(aw), Inches(0.17))
            ar.adjustments[0] = 0.42
            ar.adjustments[1] = 0.52
            ar.fill.solid(); ar.fill.fore_color.rgb = c
            ar.line.fill.background(); _noshadow(ar)
        x += step_w + gap
    return y + h


def lane_tag(s, x, y, label, icon, tone):
    fg, bg, ln = ACCENTS[tone]
    w = 0.72 + (len(label) * (10.5 * 0.60 + 1.4)) / 72.0
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.40))
    b.adjustments[0] = 0.5
    b.fill.solid(); b.fill.fore_color.rgb = bg
    b.line.color.rgb = ln; b.line.width = Pt(1.1)
    _noshadow(b)
    icon_tile(s, x + 0.07, y + 0.05, 0.30, icon, tone)
    text(s, label.upper(), x + 0.44, y + 0.09, w - 0.48, 0.24, size=10.5, color=fg, bold=True, space=1.4)
    return x + w
