"""Render a generated .pptx back to an HTML contact sheet for visual checking.

Reads the real file (geometry, fills, runs) rather than the build script, so it
shows what was actually written. Approximate typography — it proves layout,
position and overlap, not PowerPoint's own line breaking.
"""
import base64, html, os, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

EMU = 914400.0
SCALE = 96.0  # px per inch


def rgb(c):
    try:
        return "#%02X%02X%02X" % (c[0], c[1], c[2])
    except Exception:
        return None


def solid_fill(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:
            return rgb(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def line_col(sh):
    try:
        return rgb(sh.line.color.rgb)
    except Exception:
        return None


def render(path, out_html):
    prs = Presentation(path)
    sw = prs.slide_width / EMU * SCALE
    sh_ = prs.slide_height / EMU * SCALE
    parts = ["<style>body{margin:0;background:#3a3632;font-family:Calibri,'Segoe UI',Helvetica,Arial,sans-serif}"
             ".slide{position:relative;overflow:hidden;background:#FAF7F2;margin:22px auto;"
             "box-shadow:0 10px 40px rgba(0,0,0,.45)}"
             ".n{position:absolute;left:0;top:-20px;color:#bbb;font-size:13px}</style>"]
    for idx, slide in enumerate(prs.slides, 1):
        parts.append(f'<div class="slide" style="width:{sw}px;height:{sh_}px">')
        parts.append(f'<div class="n">slide {idx}</div>')
        for shp in slide.shapes:
            if shp.left is None:
                continue
            L = shp.left / EMU * SCALE
            T = shp.top / EMU * SCALE
            W = (shp.width or 0) / EMU * SCALE
            H = (shp.height or 0) / EMU * SCALE
            base = f"position:absolute;left:{L:.1f}px;top:{T:.1f}px;width:{W:.1f}px;height:{H:.1f}px;"

            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = shp.image.blob
                    ext = shp.image.ext
                    b64 = base64.b64encode(blob).decode()
                    parts.append(f'<img style="{base}object-fit:fill" src="data:image/{ext};base64,{b64}">')
                except Exception:
                    parts.append(f'<div style="{base}background:#ddd"></div>')
                continue
            if shp.shape_type == MSO_SHAPE_TYPE.MEDIA:
                parts.append(f'<div style="{base}background:#111;color:#fff;display:flex;'
                             f'align-items:center;justify-content:center;font-size:20px">▶ embedded video</div>')
                continue

            f = solid_fill(shp)
            lc = line_col(shp)
            st = base
            try:
                nm = shp._element.spPr.prstGeom.get("prst") or ""
            except Exception:
                nm = ""
            if f:
                st += f"background:{f};"
            if lc:
                st += f"border:1.2px solid {lc};box-sizing:border-box;"
            if nm == "roundRect":
                st += f"border-radius:{min(W,H)*0.16:.0f}px;"
            elif nm == "ellipse":
                st += "border-radius:50%;"
            elif nm == "rightArrow":
                st += ("clip-path:polygon(0 30%,58% 30%,58% 0,100% 50%,58% 100%,"
                       "58% 70%,0 70%);border:none;")
            parts.append(f'<div style="{st}">')
            if shp.has_text_frame:
                # honour the real vertical anchor — deckkit anchors top by
                # default, and centring everything invents overlaps that the
                # generated file does not have
                try:
                    anch = shp.text_frame.vertical_anchor
                except Exception:
                    anch = None
                just = ("center" if anch == MSO_ANCHOR.MIDDLE
                        else "flex-end" if anch == MSO_ANCHOR.BOTTOM else "flex-start")
                parts.append(f'<div style="position:absolute;inset:0;display:flex;flex-direction:column;'
                             f'justify-content:{just}">')
                for p in shp.text_frame.paragraphs:
                    al = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}.get(p.alignment, "left")
                    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
                    runs = []
                    for r in p.runs:
                        sz = r.font.size.pt if r.font.size else 18
                        col = rgb(r.font.color.rgb) if (r.font.color and r.font.color.type is not None) else "#2B2B2B"
                        try:
                            spc = r.font._rPr.get("spc")
                        except Exception:
                            spc = None
                        sp = f"letter-spacing:{int(spc)/100:.2f}pt;" if spc else ""
                        runs.append(f'<span style="font-size:{sz*SCALE/72:.1f}px;'
                                    f'font-weight:{"700" if r.font.bold else "400"};color:{col};{sp}'
                                    f'{"font-style:italic;" if r.font.italic else ""}">'
                                    f'{html.escape(r.text)}</span>')
                    parts.append(f'<div style="text-align:{al};line-height:{ls};white-space:pre-wrap">'
                                 f'{"".join(runs) or "&nbsp;"}</div>')
                parts.append("</div>")
            parts.append("</div>")
        parts.append("</div>")
    with open(out_html, "w") as fh:
        fh.write("".join(parts))
    print("wrote", out_html, f"({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
